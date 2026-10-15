# -*- coding: utf-8 -*-
"""Generate proposal_topic2.docx — A4, margins 2.5 cm, Times New Roman 12pt, single spacing."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = r"C:\Users\admin\Desktop\港城莞AI课程\1210-Artificial Intelligence(人工智能)\项目-project\proposal_topic2.docx"

doc = Document()

# Page setup: A4, 2.5 cm margins all round
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.5)

# Base style: Times New Roman 12pt, single spacing
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
style.paragraph_format.line_spacing = 1.0
style.paragraph_format.space_after = Pt(4)


def para(runs, align=None, space_after=4):
    """runs: list of (text, bold, italic) tuples."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.0
    if align:
        p.alignment = align
    for text, bold, italic in runs:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
    return p


# Title (12pt per format spec, centered, bold)
para([("Course Project Proposal", True, False),
      (" — Automated Bootstrapping Placement for Encrypted Inference", False, False)],
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

para([("Group. ", True, False),
      ("Xinghang ZENG (72656014), Chengyu BIE (72655619), Weijia ZHANG (72657103); "
       "contact email: 72656014@cityu-dg.edu.cn", False, False)])

para([("Topic and why. ", True, False),
      ("We choose Topic 2, ", False, False),
      ("Automated Bootstrapping Placement for Encrypted Inference", False, True),
      (". In CKKS-style fully homomorphic encryption, every multiplicative operator consumes a ciphertext "
       "“level”, and when levels run out an expensive Bootstrap must refresh the ciphertext. Deciding "
       "which edges of a network’s computation graph carry Bootstraps (and Modswitches) determines whether "
       "the block is level-correct and how low its estimated latency is. We choose this topic because it is a "
       "pure search-and-optimisation problem — the core of this course: it can be formulated as constrained "
       "combinatorial optimisation over a DAG, and it needs no LLM API and no GPU. Everything runs on a laptop, "
       "and progress is measurable against objective baselines.", False, False)])

para([("Tentative plan. ", True, False),
      ("Using the handout DAGs (the MLP block and the attention block, with level costs, fitted latency "
       "functions c", False, False),
      ("v", False, True), ("(l), and #cipher per edge), we decide for every edge (u,v) a chain w", False, False),
      ("u,v", False, True),
      (" ∈ {∅, B1, B2, M(ℓ), B1·M(ℓ), B2·M(ℓ)} with 0 ≤ ℓ ≤ L", False, False),
      ("max", False, True),
      (" (Bootstrap2 only on edges with #cipher 8 or 16; chains sharing a prefix insert and pay it once), "
       "minimising the estimated latency f = Σ", False, False),
      ("v", False, True), (" c", False, False), ("v", False, True), ("(l", False, False),
      ("v,in", False, True),
      (") + Σ", False, False), ("u", False, True), (" Σ", False, False),
      ("p∈Pu", False, True),
      (" lat(p), subject to the level constraints (operator level costs, Bootstrap refresh levels, compatible "
       "input levels at multi-input nodes, operator validity ranges), for every input level L", False, False),
      ("init", False, True),
      (" ∈ {0,…,14}. Our method, in stages:", False, False)])

steps = [
    ("Validation harness. ", "Reproduce the reference baselines with the handout scorer (evaluate_dag) before "
     "any optimisation, so that our reading of the rules is confirmed."),
    ("Exact search. ", "A CSP / branch-and-bound formulation over level variables and chain choices, with "
     "forward checking; expected to solve the MLP block optimally for each L_init."),
    ("Informed and local search. ", "A*-style state-space search over edge decisions with an admissible lower "
     "bound on the remaining latency, plus simulated annealing / local search with constraint-repair moves "
     "for the larger attention block."),
    ("Optional extension. ", "Solve the joint attention+MLP DAG as a single placement problem."),
]
for i, (head, body) in enumerate(steps, 1):
    p = para([(f"{i}. ", False, False), (head, True, False), (body, False, False)])
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(2)

para([("Evaluation. ", True, False),
      ("Benchmark: the two provided blocks over the full range L", False, False), ("init", False, True),
      (" = 0–14, scored by the reference scorer. Baselines to beat: the reference placements "
       "get_mlp_baseline_inserted() (≈28,130 ms, mildly dependent on L", False, False),
      ("init", False, True),
      (") and get_attn_baseline_inserted() (53,558.17 ms), and, for the extension, get_baseline() "
       "(81,688.82 ms). Success criterion: a feasible placement with strictly smaller f at every L", False, False),
      ("init", False, True),
      ("; where the exact method applies we also report the optimality gap. We will report estimated latency, "
       "#B, depth, budget and utilisation for each L", False, False), ("init", False, True),
      (", and annotated DAG figures for representative values. No LLM is used in this project.", False, False)])

doc.save(OUT)
print("saved:", OUT)
